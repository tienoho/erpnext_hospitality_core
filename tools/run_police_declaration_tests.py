"""Kịch bản kiểm thử MỚI — api/police_declaration.py (Đợt 1, mục 3 của kế
hoạch "Kịch bản kiểm thử TOÀN BỘ tính năng còn lại"). Module khai báo tạm trú
Công an/Cổng XNC Quảng Ninh đã qua NHIỀU vòng fix tĩnh (lọc sai company/customer,
thiếu gender/date_of_birth/nationality, XML injection, lệch mã giới tính
CSV/XLSX, nuốt exception) nhưng CHƯA TỪNG chạy thật lần nào (xác nhận grep 5
file test cũ — 0 kết quả). Đây là P0 pháp lý — sai sót ở đây có thể dẫn tới
báo cáo sai/thiếu cho Công an.

Chạy trên site test riêng, dữ liệu rollback sau mỗi ca.
"""
import re
import unittest
import xml.etree.ElementTree as ET

import frappe
from frappe.utils import nowdate, add_days
from run_property_integration import PropertyDatabaseTests


class PoliceDeclarationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        PropertyDatabaseTests.setUpClass()

    def tearDown(self):
        frappe.set_user('Administrator')
        frappe.db.rollback()

    def configure_police_settings(self):
        settings = frappe.get_single('Hospitality Police Settings')
        settings.establishment_name = 'Tuan Chau Resort Test'
        settings.establishment_code = 'TCR-TEST'
        settings.police_station_name = 'Cong An Phuong Tuan Chau'
        settings.police_city = 'Quang Ninh'
        settings.tax_id = '5700123456'
        settings.resort_company_name = self.company
        settings.address = '1 Tuan Chau, Ha Long'
        settings.default_stay_purpose = 'Du lịch / Nghỉ dưỡng'
        settings.save()

    def fixture(self):
        f = PropertyDatabaseTests()
        reservation, settings = f.accounting_reservation()
        self.property = reservation.property
        self.company = reservation.operating_company
        self.room_type = reservation.room_type
        self.currency = reservation.currency
        self.configure_police_settings()
        return f

    def make_reservation(self, f, guest_kwargs=None, res_kwargs=None, number='PD-1', check_in=True):
        room = f.room(self.property, self.room_type, number=number)
        guest = frappe.get_doc(dict(
            doctype='Guest', full_name=(guest_kwargs or {}).pop('full_name', 'Nguyen Van A'),
            guest_type='Regular', **(guest_kwargs or {})
        )).insert(ignore_permissions=True)
        res = frappe.get_doc(dict(
            doctype='Hotel Reservation', property=self.property, guest=guest.name, room=room.name,
            room_type=self.room_type, hotel_reception=room.hotel_reception, currency=self.currency,
            arrival_date=nowdate(), departure_date=add_days(nowdate(), 2),
            **(res_kwargs or {})
        )).insert(ignore_permissions=True)
        if check_in:
            from hospitality_core.hospitality_core.doctype.hotel_reservation.hotel_reservation import check_in_guest
            check_in_guest(res.name)
            res.reload()
        return guest, res

    # ------------------------------------------------------------------
    # 1) get_daily_guest_list(): phải bao gồm ĐỦ khách walk-in lẫn khách đặt
    #    qua công ty/đại lý (P0: filter r.company==Company cũ đã loại nhầm
    #    nhóm này — xác nhận fix còn đứng vững).
    # ------------------------------------------------------------------
    def test_daily_guest_list_includes_walkin_and_company_booked_guests(self):
        f = self.fixture()
        agency = 'Police Test Agency'
        if not frappe.db.exists('Customer', agency):
            frappe.get_doc(dict(doctype='Customer', customer_name=agency, customer_type='Company',
                customer_group=frappe.db.get_value('Customer Group', {'is_group': 0}),
                territory=frappe.db.get_value('Territory', {'is_group': 0}))).insert(ignore_permissions=True)
        _, walkin_res = self.make_reservation(f, guest_kwargs=dict(full_name='Tran Thi Walkin'), number='PD-WALKIN')
        _, company_res = self.make_reservation(f, guest_kwargs=dict(full_name='Le Van Company'),
            res_kwargs=dict(is_company_guest=1, company=agency), number='PD-COMPANY')
        from hospitality_core.hospitality_core.api.police_declaration import get_daily_guest_list
        guests = get_daily_guest_list(target_date=nowdate())
        names = {g.full_name for g in guests}
        self.assertIn('Tran Thi Walkin', names, 'Khách walk-in phải có trong danh sách khai báo.')
        self.assertIn('Le Van Company', names,
            'Khách đặt qua công ty/đại lý KHÔNG được loại khỏi danh sách khai báo tạm trú (đúng fix P0).')

    def test_daily_guest_list_excludes_not_checked_in(self):
        f = self.fixture()
        self.make_reservation(f, guest_kwargs=dict(full_name='Pham Reserved Only'), number='PD-RESV', check_in=False)
        from hospitality_core.hospitality_core.api.police_declaration import get_daily_guest_list
        guests = get_daily_guest_list(target_date=nowdate())
        names = {g.full_name for g in guests}
        self.assertNotIn('Pham Reserved Only', names, 'Đặt phòng chưa Checked In không phải khách đang lưu trú thật.')

    # ------------------------------------------------------------------
    # 2) Escape XML injection (P0 đã fix tĩnh — xác nhận thật qua parse XML).
    # ------------------------------------------------------------------
    def test_xml_export_escapes_special_characters_and_parses_valid(self):
        # Lưu ý: KHÔNG dùng '<'/'>' trong full_name ở đây — phát hiện phụ
        # (ngoài phạm vi module này): Guest.validate() tự tạo Customer mới
        # với name=customer_name=full_name khi guest chưa gắn customer, và
        # Frappe's validate_name() cấm '<'/'>' trong TÊN CHỨNG TỪ — 1 tên
        # khách chứa '<'/'>' sẽ crash NGAY LÚC TẠO GUEST (frappe.NameError),
        # trước khi tới được module khai báo công an. Rủi ro thực tế thấp
        # (tên thật hiếm khi có 2 ký tự này) nhưng là 1 giới hạn có thật của
        # `guest.py` — ghi nhận riêng, không thuộc phạm vi sửa của đợt test
        # police_declaration.py này. '&' và '"' vẫn là ký tự đặc biệt thật sự
        # cần escape cho XML và KHÔNG bị chặn ở tầng đặt tên, nên đủ để kiểm
        # tra đúng mục tiêu (XML injection) của ca test này.
        f = self.fixture()
        self.make_reservation(f, guest_kwargs=dict(full_name='Nguyen & Tran "VIP"'), number='PD-XML')
        from hospitality_core.hospitality_core.api.police_declaration import export_police_declaration_xml
        xml_content = export_police_declaration_xml(target_date=nowdate(), company=self.company)
        # Phải parse được bằng XML parser chuẩn — nếu escape sai, ET.fromstring sẽ raise.
        root = ET.fromstring(xml_content)
        names = [el.text for el in root.iter('HoTen')]
        self.assertIn('NGUYEN & TRAN "VIP"', names,
            'Tên khách có ký tự đặc biệt phải xuất hiện ĐÚNG NGUYÊN VĂN sau khi parse lại XML (escape đúng, không hỏng cấu trúc).')

    # ------------------------------------------------------------------
    # 3) Mã giới tính nhất quán CSV vs XLSX cho báo cáo XNC (P0 đã fix tĩnh).
    # ------------------------------------------------------------------
    def test_quangninh_report_gender_code_consistent_between_csv_and_xlsx(self):
        f = self.fixture()
        self.make_reservation(f, guest_kwargs=dict(full_name='Jane Foreign Smith', gender='Female',
            nationality='United Kingdom'), res_kwargs=dict(is_alien=1, passport_number='P1234567'), number='PD-FOREIGN')
        from hospitality_core.hospitality_core.api.police_declaration import (
            export_quangninh_immigration_report, export_quangninh_immigration_report_xlsx)
        csv_data = export_quangninh_immigration_report(target_date=nowdate(), company=self.company)
        self.assertIn(',2,', csv_data, 'CSV phải xuất mã giới tính "2" (Nữ) theo đúng cột chuẩn Cổng XNC.')
        import openpyxl, io
        xlsx_bytes = export_quangninh_immigration_report_xlsx(target_date=nowdate(), company=self.company)
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
        ws = wb.active
        header = [c.value for c in ws[5]]
        gender_col = next(i for i, h in enumerate(header) if 'Giới tính' in (h or ''))
        data_row = [c.value for c in ws[6]]
        self.assertEqual(str(data_row[gender_col]), '2',
            'Bản Excel phải xuất CÙNG mã giới tính "2" như bản CSV (không lệch định dạng giữa 2 bản của cùng 1 báo cáo).')

    def test_quangninh_report_only_includes_foreign_guests(self):
        f = self.fixture()
        self.make_reservation(f, guest_kwargs=dict(full_name='Vietnamese Guest Only'), number='PD-VN')
        self.make_reservation(f, guest_kwargs=dict(full_name='German Tourist', nationality='Germany'),
            res_kwargs=dict(is_alien=1, passport_number='G7654321'), number='PD-DE')
        from hospitality_core.hospitality_core.api.police_declaration import export_quangninh_immigration_report
        csv_data = export_quangninh_immigration_report(target_date=nowdate(), company=self.company)
        self.assertNotIn('GUEST ONLY', csv_data.upper(), 'Khách Việt Nam không được xuất vào báo cáo XNC người nước ngoài.')
        self.assertIn('G7654321', csv_data, 'Khách nước ngoài (có hộ chiếu) phải có mặt trong báo cáo XNC.')

    # ------------------------------------------------------------------
    # 4) Phân quyền: role ngoài ALLOWED_ROLES phải bị chặn (P0 đã fix tĩnh —
    #    bỏ điều kiện OR dự phòng làm vô hiệu ALLOWED_ROLES).
    # ------------------------------------------------------------------
    def test_unauthorized_role_blocked(self):
        self.fixture()
        user_email = 'police-unauthorized@example.com'
        if not frappe.db.exists('User', user_email):
            frappe.get_doc(dict(doctype='User', email=user_email, first_name='NoAccess',
                send_welcome_email=0, roles=[dict(role='Sales User')])).insert(ignore_permissions=True)
        frappe.set_user(user_email)
        from hospitality_core.hospitality_core.api.police_declaration import get_daily_guest_list
        with self.assertRaises(frappe.PermissionError):
            get_daily_guest_list(target_date=nowdate())

    def test_authorized_role_allowed(self):
        self.fixture()
        user_email = 'police-authorized@example.com'
        if not frappe.db.exists('User', user_email):
            frappe.get_doc(dict(doctype='User', email=user_email, first_name='Frontdesk',
                send_welcome_email=0, roles=[dict(role='Frontdesk User')])).insert(ignore_permissions=True)
        frappe.set_user(user_email)
        from hospitality_core.hospitality_core.api.police_declaration import get_daily_guest_list
        # Không được ném PermissionError — có thể trả về danh sách rỗng, không sao.
        get_daily_guest_list(target_date=nowdate())

    # ------------------------------------------------------------------
    # 5) Thiếu cấu hình bắt buộc phải báo lỗi RÕ RÀNG, không xuất báo cáo rỗng/sai.
    # ------------------------------------------------------------------
    def test_missing_required_settings_throws_clear_error(self):
        # establishment_name/establishment_code/tax_id đã được chính DocType
        # (validate_establishment_info()/validate_tax_id()) bắt buộc từ lúc
        # save() nên không bao giờ để trống được tới đây — chỉ police_station_name/
        # police_city/address KHÔNG có validate() riêng ở DocType, nên
        # _get_police_settings()'s runtime check là chốt chặn DUY NHẤT.
        self.fixture()
        settings = frappe.get_single('Hospitality Police Settings')
        settings.police_station_name = ''
        settings.save()
        from hospitality_core.hospitality_core.api.police_declaration import get_daily_guest_list
        with self.assertRaisesRegex(frappe.ValidationError, 'Tên đơn vị công an'):
            get_daily_guest_list(target_date=nowdate())

    # ------------------------------------------------------------------
    # 6) Logic thuần: _is_foreign_guest()/_normalize_iso3_nationality().
    # ------------------------------------------------------------------
    def test_foreign_guest_detection_and_iso3_normalization(self):
        from hospitality_core.hospitality_core.api.police_declaration import (
            _is_foreign_guest, _normalize_iso3_nationality)
        self.assertFalse(_is_foreign_guest(dict(nationality='Vietnam')))
        self.assertTrue(_is_foreign_guest(dict(is_alien=1)))
        self.assertTrue(_is_foreign_guest(dict(identification_type='Passport')))
        self.assertTrue(_is_foreign_guest(dict(nationality='Germany')))
        self.assertEqual(_normalize_iso3_nationality('Vietnam'), 'VNM')
        self.assertEqual(_normalize_iso3_nationality('', is_alien=True), 'FOR')
        self.assertEqual(_normalize_iso3_nationality('', is_alien=False), 'VNM')
        self.assertEqual(_normalize_iso3_nationality('DEU'), 'DEU')


if __name__ == "__main__":
    import sys, json
    try:
        result = unittest.TextTestRunner(verbosity=2).run(
            unittest.defaultTestLoader.loadTestsFromTestCase(PoliceDeclarationTests))
        print(json.dumps(dict(tests=result.testsRun, failures=len(result.failures), errors=len(result.errors))))
        sys.exit(0 if result.wasSuccessful() else 1)
    finally:
        frappe.db.rollback()
        frappe.destroy()
